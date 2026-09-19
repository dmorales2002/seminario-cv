"""
Servicio de extracción de texto desde archivos PDF y DOCX.
Compatibilidad: texto seleccionable (no imágenes escaneadas).
"""
import io
import pdfplumber
import docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extrae el texto de un PDF usando pdfplumber (maneja columnas y tablas mejor que PyPDF2)."""
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extrae el texto de un DOCX leyendo todos los párrafos del documento."""
    document = docx.Document(io.BytesIO(file_bytes))
    text_parts = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())
    # Incluir texto de tablas también
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text.strip())
    return "\n".join(text_parts).strip()


def extract_text(file_bytes: bytes, file_type: str) -> str:
    """Dispatcher principal. Retorna el texto crudo del documento."""
    if file_type == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif file_type == "docx":
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError(f"Tipo de archivo no soportado: {file_type}")
