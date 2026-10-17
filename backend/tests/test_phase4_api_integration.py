import os as _os
_os.environ.setdefault("SECRET_KEY", "clave-solo-para-pruebas-" + "x" * 32)
import io
import os
import unittest
import uuid

import docx
from fastapi.testclient import TestClient

from app.api.endpoints.analysis import get_analysis_engine
from app.core.config import settings
from app.db.session import SessionLocal
from app.main import app
from app.models.application import Application
from app.models.resume import Resume
from app.models.user import User
from app.models.vacancy import Vacancy
from app.schemas.analysis import CandidateProfile, SemanticAssessment, VacancyCriteriaDraft
from app.services.affinity import AffinityEngine


class FakeProvider:
    model = "fake-integration-model"

    def extract_vacancy_criteria(self, title, description, requirements):
        return VacancyCriteriaDraft(criteria=[
            {"name": "Python", "kind": "SKILL", "description": "Programación Python", "importance": 3},
            {"name": "Diseño de APIs", "kind": "ACTIVITY", "description": "Construcción de APIs", "importance": 1},
        ])

    def extract_candidate_profile(self, raw_text):
        return CandidateProfile(claims=[
            {"name": "Python", "kind": "SKILL", "evidence_quote": "Desarrollé servicios con Python y FastAPI.", "section": "Experiencia"},
        ])

    def match(self, criteria, profile):
        return SemanticAssessment(matches=[
            {"criterion_index": 0, "claim_index": 0, "strength": 1.0, "rationale": "Coincidencia directa."},
            {"criterion_index": 1, "claim_index": 0, "strength": 0.5, "rationale": "FastAPI evidencia parcialmente diseño de APIs."},
        ])


@unittest.skipUnless(os.getenv("RUN_DB_INTEGRATION") == "1", "requiere PostgreSQL local")
class Phase4ApiIntegrationTests(unittest.TestCase):
    def setUp(self):
        app.dependency_overrides[get_analysis_engine] = lambda: AffinityEngine(FakeProvider())
        self._saved = (settings.RECRUITER_INVITE_CODE, settings.ANALYSIS_TIME_BUDGET_SECONDS)
        settings.RECRUITER_INVITE_CODE = "codigo-de-prueba"
        self.client = TestClient(app)
        self.created_user_ids = []
        self.created_vacancy_id = None
        self.created_application_id = None

    def tearDown(self):
        app.dependency_overrides.clear()
        settings.RECRUITER_INVITE_CODE, settings.ANALYSIS_TIME_BUDGET_SECONDS = self._saved
        db = SessionLocal()
        try:
            if self.created_application_id:
                application = db.get(Application, self.created_application_id)
                if application:
                    resume = db.get(Resume, application.resume_id)
                    db.delete(application)
                    db.flush()
                    if resume:
                        db.delete(resume)
            if self.created_vacancy_id:
                for application in db.query(Application).filter(Application.vacancy_id == self.created_vacancy_id).all():
                    resume = db.get(Resume, application.resume_id)
                    db.delete(application)
                    db.flush()
                    if resume:
                        db.delete(resume)
                db.flush()
                vacancy = db.get(Vacancy, self.created_vacancy_id)
                if vacancy:
                    db.delete(vacancy)
            for user_id in self.created_user_ids:
                user = db.get(User, user_id)
                if user:
                    db.delete(user)
            db.commit()
        finally:
            db.close()

    def _register_and_login(self, role):
        email = f"phase4-{uuid.uuid4().hex}@example.com"
        password = "SecurePass123!"
        response = self.client.post("/api/v1/auth/register", json={
            "full_name": f"{role} Integration",
            "email": email,
            "password": password,
            "role": role,
            "invite_code": "codigo-de-prueba" if role == "RECRUITER" else None,
        })
        self.assertEqual(response.status_code, 200, response.text)
        self.created_user_ids.append(response.json()["id"])
        login = self.client.post("/api/v1/auth/login", data={"username": email, "password": password})
        self.assertEqual(login.status_code, 200, login.text)
        return {"Authorization": f"Bearer {login.json()['access_token']}"}

    def test_complete_analysis_flow_with_evidence_and_invalidation(self):
        recruiter_headers = self._register_and_login("RECRUITER")
        candidate_headers = self._register_and_login("CANDIDATE")

        vacancy_response = self.client.post("/api/v1/vacancies/", headers=recruiter_headers, json={
            "title": "Backend Python",
            "description": "Construir servicios y APIs.",
            "requirements": "Python, FastAPI y diseño de APIs.",
        })
        self.assertEqual(vacancy_response.status_code, 201, vacancy_response.text)
        self.created_vacancy_id = vacancy_response.json()["id"]

        document = docx.Document()
        document.add_heading("Experiencia", level=1)
        document.add_paragraph("Desarrollé servicios con Python y FastAPI.")
        document.add_paragraph("Licenciatura y 10 años de experiencia.")
        buffer = io.BytesIO()
        document.save(buffer)
        application_response = self.client.post(
            f"/api/v1/applications/{self.created_vacancy_id}",
            headers=candidate_headers,
            files={"cv_file": ("cv.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        self.assertEqual(application_response.status_code, 202, application_response.text)
        self.created_application_id = application_response.json()["id"]

        resume_response = self.client.get(
            f"/api/v1/applications/{self.created_vacancy_id}/{self.created_application_id}/resume",
            headers=recruiter_headers,
        )
        self.assertEqual(resume_response.status_code, 200, resume_response.text)
        self.assertIn("wordprocessingml.document", resume_response.headers["content-type"])
        self.assertEqual(resume_response.content, buffer.getvalue())
        self.assertEqual(resume_response.headers["cache-control"], "private, no-store")
        stranger_headers = self._register_and_login("CANDIDATE")
        forbidden = self.client.get(
            f"/api/v1/applications/{self.created_vacancy_id}/{self.created_application_id}/resume",
            headers=stranger_headers,
        )
        self.assertEqual(forbidden.status_code, 403)

        criteria = self.client.post(
            f"/api/v1/analysis/vacancies/{self.created_vacancy_id}/criteria",
            headers=recruiter_headers,
        )
        self.assertEqual(criteria.status_code, 200, criteria.text)
        self.assertEqual(sum(item["weight"] for item in criteria.json()["criteria"]), 100)

        run = self.client.post(
            f"/api/v1/analysis/vacancies/{self.created_vacancy_id}/run",
            headers=recruiter_headers,
        )
        self.assertEqual(run.status_code, 200, run.text)
        self.assertEqual(run.json()["analyzed"], 1)
        self.assertEqual(run.json()["results"][0]["score"], 88)

        repeated = self.client.post(
            f"/api/v1/analysis/vacancies/{self.created_vacancy_id}/run",
            headers=recruiter_headers,
        )
        self.assertEqual(repeated.status_code, 200, repeated.text)
        self.assertEqual(repeated.json()["results"][0]["score"], 88)
        self.assertEqual(repeated.json()["results"][0]["analyzed_at"], run.json()["results"][0]["analyzed_at"])

        ranking = self.client.get(
            f"/api/v1/applications/{self.created_vacancy_id}",
            headers=recruiter_headers,
        )
        evidence = ranking.json()[0]["ai_evidence"]
        self.assertEqual(ranking.json()[0]["analysis_status"], "COMPLETED")
        self.assertEqual(evidence["factors"][0]["evidence_quote"], "Desarrollé servicios con Python y FastAPI.")
        self.assertNotIn("Licenciatura", str(evidence))
        self.assertNotIn("10 años", str(evidence))

        edited = self.client.put(
            f"/api/v1/vacancies/{self.created_vacancy_id}",
            headers=recruiter_headers,
            json={"requirements": "Python y pruebas automatizadas."},
        )
        self.assertEqual(edited.status_code, 200, edited.text)
        invalidated = self.client.get(
            f"/api/v1/applications/{self.created_vacancy_id}",
            headers=recruiter_headers,
        ).json()[0]
        self.assertIsNone(invalidated["ai_score"])
        self.assertEqual(invalidated["analysis_status"], "PENDING")

    def test_recruiter_registration_requires_invite_code(self):
        base = {"full_name": "Intruso", "email": f"phase4-{uuid.uuid4().hex}@example.com", "password": "SecurePass123!", "role": "RECRUITER"}
        sin_codigo = self.client.post("/api/v1/auth/register", json=base)
        self.assertEqual(sin_codigo.status_code, 403)
        mal = self.client.post("/api/v1/auth/register", json={**base, "invite_code": "otro"})
        self.assertEqual(mal.status_code, 403)
        settings.RECRUITER_INVITE_CODE = None
        deshabilitado = self.client.post("/api/v1/auth/register", json={**base, "invite_code": "codigo-de-prueba"})
        self.assertEqual(deshabilitado.status_code, 403)

    def test_api_rejects_weak_passwords(self):
        for password in ("a", "corta1", "soloLetrasLargas", "123456789"):
            response = self.client.post("/api/v1/auth/register", json={
                "full_name": "Debil", "email": f"phase4-{uuid.uuid4().hex}@example.com", "password": password, "role": "CANDIDATE",
            })
            self.assertEqual(response.status_code, 422, password)

    def test_analysis_runs_in_time_limited_batches(self):
        recruiter_headers = self._register_and_login("RECRUITER")
        vacancy_response = self.client.post("/api/v1/vacancies/", headers=recruiter_headers, json={
            "title": "Backend Python", "description": "Construir servicios.", "requirements": "Python y FastAPI.",
        })
        self.created_vacancy_id = vacancy_response.json()["id"]
        document = docx.Document()
        document.add_paragraph("Desarrollé servicios con Python y FastAPI.")
        buffer = io.BytesIO()
        document.save(buffer)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        for _ in range(3):
            headers = self._register_and_login("CANDIDATE")
            response = self.client.post(f"/api/v1/applications/{self.created_vacancy_id}", headers=headers,
                                        files={"cv_file": ("cv.docx", buffer.getvalue(), mime)})
            self.assertEqual(response.status_code, 202, response.text)
        settings.ANALYSIS_TIME_BUDGET_SECONDS = 0
        pendientes = []
        for _ in range(5):
            run = self.client.post(f"/api/v1/analysis/vacancies/{self.created_vacancy_id}/run", headers=recruiter_headers)
            self.assertEqual(run.status_code, 200, run.text)
            pendientes.append(run.json()["pending"])
            if run.json()["complete"]:
                break
        self.assertEqual(pendientes, [2, 1, 0])
        ranking = self.client.get(f"/api/v1/applications/{self.created_vacancy_id}", headers=recruiter_headers).json()
        self.assertEqual([item["analysis_status"] for item in ranking], ["COMPLETED"] * 3)


if __name__ == "__main__":
    unittest.main()
