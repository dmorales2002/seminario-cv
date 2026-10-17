import os as _os
_os.environ.setdefault("SECRET_KEY", "clave-solo-para-pruebas-" + "x" * 32)
import io
import unittest

import docx
from fastapi import HTTPException
from pydantic import ValidationError

from app.api.endpoints.applications import validate_cv_content
from app.models.vacancy import VacancyStatus
from app.schemas.vacancy import VacancyCreate, VacancyUpdate
from app.services.extractor import extract_text, extract_text_from_docx


class VacancySchemaTests(unittest.TestCase):
    def test_create_rejects_whitespace_only_fields(self):
        with self.assertRaises(ValidationError):
            VacancyCreate(title="  ", description="Descripción", requirements="Python")

    def test_update_accepts_paused_status(self):
        update = VacancyUpdate(status=VacancyStatus.PAUSED)
        self.assertEqual(update.status, VacancyStatus.PAUSED)


class CvValidationTests(unittest.TestCase):
    def test_rejects_empty_file(self):
        with self.assertRaises(HTTPException) as context:
            validate_cv_content(b"", "pdf")
        self.assertEqual(context.exception.status_code, 400)

    def test_rejects_spoofed_pdf(self):
        with self.assertRaises(HTTPException) as context:
            validate_cv_content(b"not really a pdf", "pdf")
        self.assertEqual(context.exception.status_code, 415)

    def test_accepts_pdf_signature(self):
        validate_cv_content(b"%PDF-1.7\n", "pdf")

    def test_extracts_docx_paragraphs_and_tables(self):
        document = docx.Document()
        document.add_paragraph("Python y FastAPI")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "PostgreSQL"
        buffer = io.BytesIO()
        document.save(buffer)

        text = extract_text_from_docx(buffer.getvalue())

        self.assertIn("Python y FastAPI", text)
        self.assertIn("PostgreSQL", text)

    def test_rejects_unsupported_document_type(self):
        with self.assertRaises(ValueError):
            extract_text(b"plain text", "txt")


if __name__ == "__main__":
    unittest.main()


class NormalizeMatchesTests(unittest.TestCase):
    def test_uncited_credit_becomes_no_evidence(self):
        from app.services.llm_client import normalize_matches
        data = {"matches": [
            {"criterion_index": 0, "claim_index": None, "strength": 0.5, "rationale": "x"},
            {"criterion_index": 1, "claim_index": 2, "strength": 0.0, "rationale": "y"},
            {"criterion_index": 2, "claim_index": 1, "strength": 1.0, "rationale": "z"},
        ]}
        out = normalize_matches(data)["matches"]
        self.assertEqual((out[0]["strength"], out[0]["claim_index"]), (0.0, None))
        self.assertIsNone(out[1]["claim_index"])
        self.assertEqual((out[2]["strength"], out[2]["claim_index"]), (1.0, 1))
